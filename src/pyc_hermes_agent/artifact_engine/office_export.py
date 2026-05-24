"""Structured PPTX/XLSX export (requires optional ``office`` deps)."""

from __future__ import annotations

from io import BytesIO
from typing import Any


def export_xlsx_bytes(spec: dict[str, Any]) -> bytes:
    """Build ``.xlsx`` from ``{"sheets": [{"name":"S","rows":[["h1"],["cell"]]}]}``."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ValueError("XLSX export requires ``openpyxl``. Install extras: ``uv sync --extra office``.") from exc

    wb = Workbook()
    sheets = spec.get("sheets")
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("spec['sheets'] must be a non-empty list")

    wb.remove(wb.active)

    for sh in sheets:
        if not isinstance(sh, dict):
            raise ValueError("Each sheet entry must be an object")

        title = str(sh.get("name", "Sheet") or "Sheet")[:31]
        ws = wb.create_sheet(title=title)

        rows = sh.get("rows") or []
        if isinstance(rows, list):
            for r_idx, row in enumerate(rows, start=1):
                if not isinstance(row, (list, tuple)):
                    ws.cell(row=r_idx, column=1, value=str(row))
                    continue
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def export_pptx_bytes(spec: dict[str, Any]) -> bytes:
    """
    Minimal deck: ``{"title":"Deck","slides":[{"title":"Slide1","bullets":["a"]}]}``
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise ValueError("PPTX export requires ``python-pptx``. Install extras: ``uv sync --extra office``.") from exc

    prs = Presentation()

    slides = spec.get("slides")
    deck_title = str(spec.get("title", "Exported"))[:200]

    if not isinstance(slides, list) or not slides:
        slides = [{"title": deck_title, "bullets": ["(empty)"]}]

    for idx, slide_spec in enumerate(slides):
        if not isinstance(slide_spec, dict):
            raise ValueError("Each slide must be an object")

        slide_title = str(slide_spec.get("title", f"Slide {idx + 1}"))[:120]

        body = slide_spec.get("bullets") or []
        if isinstance(body, list):
            bullets = body
        elif isinstance(body, str):
            bullets = [body]
        else:
            bullets = []

        layout_idx = 6 if len(prs.slide_layouts) > 6 else min(5, len(prs.slide_layouts) - 1)
        layout = prs.slide_layouts[layout_idx]
        slide = prs.slides.add_slide(layout)

        left = Inches(0.35)
        top = Inches(0.25)
        tb = slide.shapes.add_textbox(left, top, Inches(9), Inches(5.8))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = slide_title
        title_para = tf.paragraphs[0]
        if title_para.runs:
            title_para.runs[0].font.size = Pt(24)

        for line in bullets:
            bp = tf.add_paragraph()
            bp.text = str(line)
            bp.level = 0
            if bp.runs:
                bp.runs[0].font.size = Pt(16)

    bio = BytesIO()
    prs.save(bio)
    return bio.getvalue()


__all__ = ["export_pptx_bytes", "export_xlsx_bytes"]
